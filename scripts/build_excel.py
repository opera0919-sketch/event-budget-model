"""엑셀 워크북 생성: 데이터셋 요약 / 케이스 비교 / 목표별 최적안 / Pareto(차트)."""
from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from src.cases import reference_cases  # noqa: E402
from src.data_generator import generate_all  # noqa: E402
from src.optimizer import optimize  # noqa: E402
from src.reward_engine import CURRENT_STRUCTURE  # noqa: E402
from src.simulation import build_cache, evaluate_structure  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RESULTS_DIR = os.path.join(ROOT, "results")


def _autosize(ws):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 28)


def _write_budget_sheet(wb, datasets, header_font, header_fill) -> None:
    """예산집행표: 추천안별·구간별 인원/리워드/소요예산 (시나리오 분리).

    탄력성을 반영하지 않은 **보수적 집계**(신청자 유지 가정)다. 예산 승인은
    이 수치와 worst-case로 요청하는 것이 안전하다.
    """
    from openpyxl.styles import Font

    from config import simulation_config as C
    from src.cases import recommended_plans
    from src.reward_engine import budget_amount

    ws = wb.create_sheet("예산집행표")
    bold = Font(bold=True)
    edges = list(C.DIRECT_BRACKETS)
    uppers = edges[1:] + [None]

    scenarios = [("전체(평균)", None), ("흥행", "hit"), ("비흥행", "flop")]
    row = 1
    for plan in recommended_plans():
        ws.cell(row=row, column=1, value=f"[{plan.name}]").font = bold
        row += 1
        for label, tag in scenarios:
            subset = [d for d in datasets if tag is None or d.scenario == tag]
            n_rounds = len(subset)
            transfers = [t for d in subset for t in d.transfers]

            ws.cell(row=row, column=1, value=f"시나리오: {label} (회차 {n_rounds})").font = bold
            row += 1
            headers = ["순입금 구간", "인원(회차평균)", "리워드", "1인예산(제세포함)",
                       "소요예산", "예산비중"]
            for c, h in enumerate(headers, start=1):
                cell = ws.cell(row=row, column=c, value=h)
                cell.font = header_font
                cell.fill = header_fill
            row += 1

            start = row
            total_budget = 0.0
            data = []
            for i, lo in enumerate(edges):
                hi = uppers[i]
                group = [t for t in transfers
                         if lo <= t < (hi if hi is not None else 10 ** 15)]
                if not group:
                    continue
                reward = plan.rewards()[i]
                per = budget_amount(reward)
                cnt = len(group) / n_rounds
                budget = cnt * per
                total_budget += budget
                label_txt = f"{lo:,}원 이상" if hi is None else f"{lo:,}~{hi:,}원"
                data.append((label_txt, cnt, reward, per, budget))
            for label_txt, cnt, reward, per, budget in data:
                ws.cell(row=row, column=1, value=label_txt)
                ws.cell(row=row, column=2, value=round(cnt, 1))
                ws.cell(row=row, column=3, value=reward)
                ws.cell(row=row, column=4, value=round(per))
                ws.cell(row=row, column=5, value=round(budget))
                ws.cell(row=row, column=6,
                        value=round(budget / total_budget, 4) if total_budget else 0)
                ws.cell(row=row, column=6).number_format = "0.0%"
                row += 1
            # 합계
            ws.cell(row=row, column=1, value="합계").font = bold
            ws.cell(row=row, column=2, value=f"=SUM(B{start}:B{row-1})").font = bold
            ws.cell(row=row, column=5, value=f"=SUM(E{start}:E{row-1})").font = bold
            row += 2
        row += 1

    ws.cell(row=row, column=1,
            value="주: 탄력성 미반영(신청자 유지 가정) 보수적 집계. "
                  "모델 추정 예산은 '케이스비교' 시트 참조. "
                  "예산 확보는 worst-case 기준 권고.")
    _autosize(ws)


def main() -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.chart import Reference, ScatterChart, Series
        from openpyxl.styles import Font, PatternFill
    except Exception:
        print("openpyxl 미설치 - CSV 산출물(results/*.csv)로 대체하세요.")
        return

    datasets = generate_all()
    caches = [build_cache(d) for d in datasets]
    out = optimize(caches)
    case_metrics = [evaluate_structure(caches, s) for s in reference_cases()]

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2C6FBB")

    def write_sheet(title, headers, rows):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.font = header_font
            c.fill = header_fill
        for r in rows:
            ws.append(r)
        _autosize(ws)
        return ws

    wb.remove(wb.active)

    # 1) 데이터셋 요약
    ds_rows = []
    for d in datasets:
        recipients = sum(1 for b in d.benefits if b > 0)
        ds_rows.append([d.round_idx, d.scenario, d.n_customers, recipients,
                        sum(d.transfers), sum(d.benefits)])
    write_sheet("데이터셋요약", ["회차", "시나리오", "신청자수", "수령자수", "이전금액합", "현행혜택합"], ds_rows)

    # 2) 케이스 비교
    keys = list(case_metrics[0].as_row().keys())
    write_sheet("케이스비교", keys, [list(m.as_row().values()) for m in case_metrics])

    # 3) 목표별 최적안
    out.baseline.name = "현행"
    names = {"A_min_budget": "A_예산최소", "B_max_efficiency": "B_효율최대", "C_target_budget": "C_목표예산"}
    obj_rows = [list(out.baseline.as_row().values())]
    for k, nm in names.items():
        out.objectives[k].metrics.name = nm
        obj_rows.append(list(out.objectives[k].metrics.as_row().values()))
    ws_obj = write_sheet("목표별최적안", keys, obj_rows)
    # 추천 구조(리워드 벡터) 부기
    ws_obj.append([])
    ws_obj.append(["구조", "배수", "리워드벡터(인정금액 티어 순)"])
    ws_obj.append(["현행", CURRENT_STRUCTURE.multiplier,
                   " / ".join(str(r) for r in CURRENT_STRUCTURE.rewards())])
    for k, nm in names.items():
        s = out.objectives[k].structure
        ws_obj.append([nm, s.multiplier, " / ".join(str(r) for r in s.rewards())])

    # 3-1) 예산집행표 (구간별 × 시나리오)
    _write_budget_sheet(wb, datasets, header_font, header_fill)

    # 4) Pareto + 차트
    ws_p = wb.create_sheet("Pareto")
    ws_p.append(["예산_평균(원)", "유치이전금액_평균(원)", "효율_배", "매력도지수"])
    for c in ws_p[1]:
        c.font = header_font
        c.fill = header_fill
    pareto_sorted = sorted(out.pareto, key=lambda sm: sm[1].budget_mean)
    for s, m in pareto_sorted:
        ws_p.append([round(m.budget_mean), round(m.transfer_mean), round(m.efficiency, 1),
                     round(m.attractiveness_index, 4)])
    _autosize(ws_p)

    n = len(pareto_sorted)
    if n >= 2:
        chart = ScatterChart()
        chart.title = "예산 vs 유치 이전금액 (Pareto Frontier)"
        chart.x_axis.title = "예산(원)"
        chart.y_axis.title = "유치 이전금액(원)"
        chart.style = 13
        xref = Reference(ws_p, min_col=1, min_row=2, max_row=n + 1)
        yref = Reference(ws_p, min_col=2, min_row=1, max_row=n + 1)
        series = Series(yref, xref, title_from_data=True)
        series.marker.symbol = "circle"
        series.graphicalProperties.line.noFill = True
        chart.series.append(series)
        chart.height = 11
        chart.width = 20
        ws_p.add_chart(chart, "F2")

    # PNG 임베드(있으면)
    png_path = os.path.join(RESULTS_DIR, "pareto.png")
    if os.path.exists(png_path):
        try:
            from openpyxl.drawing.image import Image as XLImage
            ws_img = wb.create_sheet("Pareto차트(PNG)")
            ws_img.add_image(XLImage(png_path), "A1")
        except Exception:
            pass

    os.makedirs(RESULTS_DIR, exist_ok=True)
    xlsx_path = os.path.join(RESULTS_DIR, "simulation_workbook.xlsx")
    wb.save(xlsx_path)
    print("엑셀 워크북 생성:", xlsx_path)


if __name__ == "__main__":
    main()
